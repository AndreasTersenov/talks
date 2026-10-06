"""Turn harvested slide states into a .pptx of native objects.

usage: build.py HARVEST_DIR OUT.pptx
"""
import glob
import json
import os
import sys
import urllib.parse

import numpy as np
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

REPO = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))
HARVEST, OUTFILE = sys.argv[1], sys.argv[2]
POSTERS = os.path.join(HARVEST, 'posters')     # optional fallback posters (qlmanage -t) for videos Chrome cannot decode
PX = 9525                      # EMU per viewport CSS px: 1280 px -> 13.333 in
E = lambda v: Emu(int(round(v * PX)))
ALIGN = {'left': PP_ALIGN.LEFT, 'start': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT,
         'end': PP_ALIGN.RIGHT, 'justify': PP_ALIGN.LEFT, '-webkit-center': PP_ALIGN.CENTER}
FONT_MAP = {'KaTeX_Main': 'Times New Roman', 'KaTeX_Math': 'Times New Roman', 'system-ui': 'IBM Plex Sans',
            'sans-serif': 'IBM Plex Sans', 'serif': 'Source Serif 4', 'monospace': 'IBM Plex Mono',
            'Source Serif 4 Display': 'Source Serif 4'}

# legacy (nameID 1) family names, which is how PowerPoint and Keynote select these weights
SEMIBOLD = {'Source Serif 4': 'Source Serif 4 Semibold', 'IBM Plex Sans': 'IBM Plex Sans SmBld'}
MEDIUM = {'Source Serif 4': 'Source Serif 4', 'IBM Plex Sans': 'IBM Plex Sans Medm', 'IBM Plex Mono': 'IBM Plex Mono Medm'}



def serif_optical(css_px):
    # Chrome sets the variable Source Serif 4 at opsz = the CSS pixel size; the static release ships
    # the named masters as separate families (Caption 8, SmText 16, Text 20, Subhead 32, Display 60)
    if css_px < 12:
        return 'Source Serif 4 Caption'
    if css_px < 18:
        return 'Source Serif 4 SmText'
    if css_px < 26:
        return 'Source Serif 4'
    if css_px < 46:
        return 'Source Serif 4 Subhead'
    return 'Source Serif 4 Display'


def hexcol(c):
    return '%02X%02X%02X' % tuple(max(0, min(255, int(round(v)))) for v in c[:3])


def set_alpha(fill_parent, alpha):
    """add <a:alpha> to the srgbClr inside an element (a solidFill)"""
    if alpha >= 0.995:
        return
    clr = fill_parent.find('.//' + qn('a:srgbClr'))
    if clr is not None:
        a = etree.SubElement(clr, qn('a:alpha'))
        a.set('val', str(int(round(max(0.0, alpha) * 100000))))


def local_path(url):
    path = urllib.parse.unquote(urllib.parse.urlparse(url).path)
    return os.path.join(REPO, path.lstrip('/'))


def picture_file(url):
    p = local_path(url)
    if p.lower().endswith('.webp'):
        png = os.path.join(HARVEST, 'conv_' + os.path.basename(p) + '.png')
        if not os.path.exists(png):
            Image.open(p).save(png)
        return png
    return p


def pic_alpha(pic, op):
    if op >= 0.995:
        return
    blip = pic._element.find('.//' + qn('a:blip'))
    a = etree.SubElement(blip, qn('a:alphaModFix'))
    a.set('amt', str(int(round(op * 100000))))


def describe(shape, text):
    if text:
        shape._element.find('.//' + qn('p:cNvPr')).set('descr', text[:1000])


def add_box(slide, it):
    r = it['rect']
    rounded = it['radius'] > 0.5 and not it.get('ellipse')
    kind = MSO_SHAPE.OVAL if it.get('ellipse') else MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(kind, E(r['x']), E(r['y']), E(r['w']), E(r['h']))
    if rounded:
        shp.adjustments[0] = min(0.5, it['radius'] / max(1, min(r['w'], r['h'])))
    shp.shadow.inherit = False
    if it['fill']:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(hexcol(it['fill']))
        set_alpha(shp.fill._xPr.find(qn('a:solidFill')), it['fill'][3] * it['op'])
    else:
        shp.fill.background()
    b = it['border']
    if b:
        shp.line.color.rgb = RGBColor.from_string(hexcol(b['c']))
        shp.line.width = Pt(b['w'] * 0.75)
        set_alpha(shp.line._get_or_add_ln().find(qn('a:solidFill')), b['c'][3] * it['op'])
        if b['dash']:
            shp.line.dash_style = 4          # MSO_LINE.DASH
    else:
        shp.line.fill.background()
    shp.text_frame.text = ''
    sh = it.get('shadow')
    if isinstance(sh, dict):
        # a CSS drop shadow as an outer shadow: blur, distance and direction from the offsets
        sppr = shp._element.spPr
        eff = sppr.find(qn('a:effectLst'))
        if eff is None:
            eff = etree.SubElement(sppr, qn('a:effectLst'))
        o = etree.SubElement(eff, qn('a:outerShdw'))
        dist = (sh['dx'] ** 2 + sh['dy'] ** 2) ** 0.5
        o.set('blurRad', str(int(sh['blur'] * PX)))
        o.set('dist', str(int(dist * PX)))
        o.set('dir', str(int((np.degrees(np.arctan2(sh['dy'], sh['dx'])) % 360) * 60000)))
        o.set('algn', 'ctr'); o.set('rotWithShape', '0')
        c = etree.SubElement(o, qn('a:srgbClr')); c.set('val', hexcol(sh['c']))
        etree.SubElement(c, qn('a:alpha')).set('val', str(int(sh['c'][3] * it['op'] * 100000)))
    return shp


def add_line(slide, it):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(it['x1']), E(it['y1']), E(it['x2']), E(it['y2']))
    ln.line.color.rgb = RGBColor.from_string(hexcol(it['c']))
    ln.line.width = Pt(it['w'] * 0.75)
    set_alpha(ln.line._get_or_add_ln().find(qn('a:solidFill')), it['c'][3] * it['op'])
    if it['dash']:
        ln.line.dash_style = 4
    ln.shadow.inherit = False


def style_run(run, r):
    f = run.font
    name = FONT_MAP.get(r['font'], r['font'])
    if name == 'Source Serif 4' and r.get('cssPx'):
        name = serif_optical(r['cssPx'])
    weight = r.get('weight', 700 if r['bold'] else 400)
    bold = weight >= 700
    if 550 <= weight < 700 and name.startswith('Source Serif 4'):
        name = name + ' Semibold'
    elif 550 <= weight < 700:
        if name in SEMIBOLD:
            name = SEMIBOLD[name]
        else:
            bold = True
    elif 450 <= weight < 550 and name in MEDIUM:
        name = MEDIUM[name]
    f.name = name
    f.size = Pt(r['pt'])
    # only ever write b="1": an explicit b="0" makes Keynote drop a weight carried by the family name
    f.bold = True if bold else None
    f.italic = r['italic']
    f.underline = r['underline'] or None
    f.color.rgb = RGBColor.from_string(hexcol(r['color']))
    rpr = run._r.get_or_add_rPr()
    set_alpha(rpr.find(qn('a:solidFill')), r['color'][3])
    if abs(r.get('spacingPt', 0)) > 0.05:
        rpr.set('spc', str(int(round(r['spacingPt'] * 100))))
    if r['strike']:
        rpr.set('strike', 'sngStrike')
    if r['shift']:
        rpr.set('baseline', '30000' if r['shift'] > 0 else '-25000')


def add_text(slide, it):
    r, c = it['rect'], it['content']
    multiline = it['lines'] > 1
    align = it['align']
    # geometry: wrapping text gets its element's content box, so it breaks where the browser broke it;
    # a single line is not allowed to wrap at all, anchored on the side its alignment anchors
    if multiline:
        x, w = c['x'], c['w'] * 1.02
    elif align in ('center', '-webkit-center'):
        x, w = c['x'], c['w']
    elif align in ('right', 'end'):
        x, w = c['x'], c['w']
    else:
        x, w = r['x'], max(r['w'] * 1.2 + 20, 4)     # slack: Keynote shrinks a no-wrap line that does not fit its box
    first_pt = next((q['pt'] for q in it['runs'] if not q.get('br')), 12)
    indent = 0
    if it['bullet']:
        indent = first_pt / 0.75 * 1.15       # px, about where the browser hangs its marker
        x -= indent
        w += indent
    box = slide.shapes.add_textbox(E(x), E(r['y']), E(w), E(r['h']))
    tf = box.text_frame
    tf.word_wrap = multiline          # a single line never wraps, whatever the font metrics do
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP

    paragraphs = [[]]
    for q in it['runs']:
        if q.get('br'):
            paragraphs.append([])
        else:
            paragraphs[-1].append(q)
    lh = it['lh'] or (r['h'] / max(1, it['lines']) if multiline else None)
    for i, runs in enumerate(paragraphs):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = ALIGN.get(align, PP_ALIGN.LEFT)
        if lh:
            para.line_spacing = Pt(lh * 0.75)
        para.space_before = para.space_after = 0
        if i == 0 and not it['bullet'] and multiline and it.get('firstIndent', 0) > 3 and align in ('left', 'start', 'justify'):
            para._p.get_or_add_pPr().set('indent', str(int(it['firstIndent'] * PX)))
        if it['bullet'] and i == 0:
            b = it['bullet']
            ppr = para._p.get_or_add_pPr()
            ppr.set('marL', str(int(indent * PX)))
            ppr.set('indent', str(-int(indent * PX)))
            clr = etree.SubElement(ppr, qn('a:buClr'))
            etree.SubElement(clr, qn('a:srgbClr')).set('val', hexcol(b['color']))
            if b.get('num'):
                etree.SubElement(ppr, qn('a:buFont')).set('typeface', '+mj-lt')
                an = etree.SubElement(ppr, qn('a:buAutoNum'))
                an.set('type', 'arabicPeriod'); an.set('startAt', str(b['num']))
            else:
                etree.SubElement(ppr, qn('a:buFont')).set('typeface', 'Arial')
                etree.SubElement(ppr, qn('a:buChar')).set('char', b['char'])
        elif it['bullet']:
            ppr = para._p.get_or_add_pPr()
            ppr.set('marL', str(int(indent * PX)))
        for q in runs:
            run = para.add_run()
            text = q['t'].replace('∘', '°') if q.get('shift', 0) > 0 else q['t']
            if q.get('math') and text.strip() in ('±', '∓'):
                text = ' ' + text.strip() + ' '
            run.text = text
            style_run(run, q)
    return box


def add_img(slide, it):
    full, clip = it['rect'], it['clip']
    pic = slide.shapes.add_picture(it.get('path') or picture_file(it['src']), E(clip['x']), E(clip['y']), E(clip['w']), E(clip['h']))
    if full['w'] > 0 and full['h'] > 0:
        pic.crop_left = (clip['x'] - full['x']) / full['w']
        pic.crop_top = (clip['y'] - full['y']) / full['h']
        pic.crop_right = (full['x'] + full['w'] - clip['x'] - clip['w']) / full['w']
        pic.crop_bottom = (full['y'] + full['h'] - clip['y'] - clip['h']) / full['h']
    pic_alpha(pic, it['op'])
    describe(pic, it.get('alt'))


def add_snap(slide, it):
    r = it['rect']
    pic = slide.shapes.add_picture(os.path.join(HARVEST, it['file']), E(r['x']), E(r['y']), E(r['w']), E(r['h']))
    # no opacity here: the snapshot was rendered with its fade already applied
    describe(pic, ('LaTeX: ' + it['alt']) if it.get('alt') and '\\' in it['alt'] else it.get('alt'))


def add_canvas(slide, it, key, n):
    import base64
    path = os.path.join(HARVEST, f'{key}_canvas{n}.png')
    with open(path, 'wb') as fh:
        fh.write(base64.b64decode(it['data'].split(',', 1)[1]))
    add_img(slide, {'rect': it['rect'], 'clip': it['clip'], 'op': it['op'], 'alt': '', 'src': None, 'path': path})


def add_video(slide, it):
    r = it['rect']
    src = local_path(it['src'])
    poster = os.path.join(POSTERS, os.path.basename(src) + '.png')
    ref = Image.open(os.path.join(HARVEST, CURRENT_KEY + '_ref.png')).convert('RGB')
    s = ref.width / 1280
    crop = ref.crop((int(r['x'] * s), int(r['y'] * s), int((r['x'] + r['w']) * s), int((r['y'] + r['h']) * s)))
    if np.asarray(crop).mean() > 8:
        poster = os.path.join(HARVEST, CURRENT_KEY + '_poster.png')
        crop.save(poster)
    slide.shapes.add_movie(src, E(r['x']), E(r['y']), E(r['w']), E(r['h']),
                           poster_frame_image=poster if os.path.exists(poster) else None, mime_type='video/mp4')


def set_background(slide, prs, key):
    bg = Image.open(os.path.join(HARVEST, key + '_bg.png')).convert('RGB')
    a = np.asarray(bg.resize((160, 90)), dtype=np.int16)
    if (a.max(axis=(0, 1)) - a.min(axis=(0, 1))).max() <= 6:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(hexcol(a.reshape(-1, 3).mean(0)))
        return
    jpg = os.path.join(HARVEST, key + '_bg.jpg')
    bg.save(jpg, quality=90)
    pic = slide.shapes.add_picture(jpg, 0, 0, prs.slide_width, prs.slide_height)
    describe(pic, 'slide background')


prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
layout = prs.slide_layouts[6]
ADD = {'box': add_box, 'line': add_line, 'text': add_text, 'img': add_img, 'snap': add_snap, 'video': add_video}
files = sorted(glob.glob(os.path.join(HARVEST, '*.json')), key=lambda f: [int(x) for x in os.path.basename(f)[:-5].split('_')])
for fn in files:
    data = json.load(open(fn))
    CURRENT_KEY = data['key']
    slide = prs.slides.add_slide(layout)
    set_background(slide, prs, data['key'])
    for n, it in enumerate(data['items']):
        if it['type'] == 'canvas':
            add_canvas(slide, it, data['key'], n)
        else:
            ADD[it['type']](slide, it)
prs.save(OUTFILE)
print(len(files), 'slides ->', OUTFILE, os.path.getsize(OUTFILE) // 2**20, 'MB')
